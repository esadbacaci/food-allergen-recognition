# -*- coding: utf-8 -*-
"""Blind expert (dietitian) audit form for the allergen database.

    python scripts/make_expert_audit_form.py

Sample: 60 dish classes, stratified and drawn with a fixed seed - 28 Turkish and 28 Food-101
composite dishes plus 2 + 2 single-ingredient classes. The form shows the dish name and one
photograph but NOT the database entry (blind). The expert marks the 14 regulated allergens
(EU 1169/2011 Annex II = TGK Annex 1) as always / often / possible / absent.

Outputs: data/allergen_db_v2/expert_audit/sample.csv
         data/allergen_db_v2/expert_audit/Alerjen_Uzman_Degerlendirme_Formu.xlsx
Scored later by scripts/score_expert_audit.py (precision, recall, Cohen's kappa per allergen).
"""
import csv
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from foodallergen import paths  # noqa: E402
from foodallergen.data import read_class_map, read_manifest  # noqa: E402

OUT = paths.REPO_DATA / "allergen_db_v2" / "expert_audit"
SEED = 42
ALLERGENS = [("GLU", "Gluten içeren tahıllar"), ("CRU", "Kabuklular"), ("EGG", "Yumurta"),
             ("FSH", "Balık"), ("PNT", "Yer fıstığı"), ("SOY", "Soya"), ("MLK", "Süt"),
             ("NUT", "Sert kabuklu meyveler"), ("CEL", "Kereviz"), ("MUS", "Hardal"),
             ("SES", "Susam"), ("SUL", "Sülfitler (>10 mg/kg)"), ("LUP", "Acı bakla (lupin)"),
             ("MOL", "Yumuşakçalar")]
CHOICES = ["Her zaman", "Sıklıkla", "Olabilir", "Yok"]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    classes = read_class_map()
    rng = random.Random(SEED)
    pick = []
    for group in ("turkish", "food101"):
        comp = sorted(c["class_id"] for c in classes if c["group"] == group and c["raw_produce"] != "1")
        simple = sorted(c["class_id"] for c in classes if c["group"] == group and c["raw_produce"] == "1")
        if not simple:  # Food-101 has no raw-produce class: take simple dishes instead
            simple = sorted(set(["edamame", "steak", "oysters", "french-fries"]) & set(comp))
            comp = [c for c in comp if c not in simple]
        pick += rng.sample(comp, 28) + rng.sample(simple, 2)
    rng.shuffle(pick)
    by_id = {c["class_id"]: c for c in classes}

    rows = read_manifest()
    photo = {}
    for r in rows:  # one training photograph per class (fixed: lowest uid)
        if r["split"] == "train" and (r["cls"] not in photo or r["uid"] < photo[r["cls"]]):
            photo[r["cls"]] = r["uid"]
    path_of = {(r["cls"], r["uid"]): r["path"] for r in rows}

    with open(OUT / "sample.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["no", "class_id", "display_tr", "display_en", "group", "photo"])
        for i, c in enumerate(pick, 1):
            w.writerow([i, c, by_id[c]["display_tr"], by_id[c]["display_en"], by_id[c]["group"],
                        path_of[(c, photo[c])]])

    wb = Workbook()
    ins = wb.active
    ins.title = "Yönerge"
    text = [
        "Alerjen veri tabanı – uzman (diyetisyen) değerlendirme formu",
        "",
        "Amaç: Bir yemek tanıma sisteminin kullandığı alerjen veri tabanının doğruluğunu bağımsız olarak ölçmek.",
        "Lütfen her yemeği, Türkiye'de / restoranlarda YAYGIN hazırlanan standart tarifine göre değerlendirin.",
        "Fotoğraf yalnızca yemeği tanımlamak içindir; fotoğraftaki tek bir örneği değil, yemeği genel olarak düşünün.",
        "",
        "Her alerjen için bir seçenek seçin:",
        "  Her zaman – standart tarifin parçası; hemen her hazırlanışta bulunur",
        "  Sıklıkla  – yaygın tariflerin önemli bir kısmında bulunur",
        "  Olabilir  – tarife bağlı, gizli içerik, garnitür ya da tipik çapraz bulaşma",
        "  Yok       – beklenmez",
        "",
        "Alerjen listesi: AB 1169/2011 Ek II ve TGK Gıda Etiketleme ve Tüketicileri Bilgilendirme Yönetmeliği Ek-1'deki 14 alerjen.",
        "Rafine yağlar ve yüksek derecede rafine nişastalar dikkate alınmaz.",
        "Lütfen değerlendirme sırasında başka bir kaynağa ya da veri tabanına bakmayın; bilgi ve deneyiminize göre yanıtlayın.",
        "Tahmini süre: 60 yemek × ~1,5 dk ≈ 1,5 saat.",
        "",
        "Değerlendiren: ____________________   Unvan: ____________________   Tarih: ____________",
    ]
    for i, t in enumerate(text, 1):
        ins.cell(row=i, column=1, value=t)
    ins["A1"].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Form")
    head = ["No", "Yemek (TR)", "Dish (EN)", "Fotoğraf"] + [f"{n}\n({c})" for c, n in ALLERGENS] + ["Not"]
    ws.append(head)
    for col in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.row_dimensions[1].height = 48
    dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    thumbs = OUT / "_thumbs"
    thumbs.mkdir(exist_ok=True)
    for i, c in enumerate(pick, 2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=by_id[c]["display_tr"] or by_id[c]["display_en"])
        ws.cell(row=i, column=3, value=by_id[c]["display_en"])
        src = paths.V2_IMAGES / path_of[(c, photo[c])]
        th = thumbs / f"{c}.png"
        with Image.open(src) as im:
            im = im.convert("RGB")
            im.thumbnail((110, 110))
            im.save(th)
        img = XLImage(str(th))
        ws.add_image(img, f"D{i}")
        ws.row_dimensions[i].height = 86
        dv.add(f"E{i}:{chr(ord('E') + len(ALLERGENS) - 1)}{i}")
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 17
    for k in range(len(ALLERGENS)):
        ws.column_dimensions[chr(ord("E") + k)].width = 12
    ws.column_dimensions[chr(ord("E") + len(ALLERGENS))].width = 30
    ws.freeze_panes = "E2"
    out = OUT / "Alerjen_Uzman_Degerlendirme_Formu.xlsx"
    wb.save(out)
    print("form:", out, "| dishes:", len(pick))


if __name__ == "__main__":
    main()
