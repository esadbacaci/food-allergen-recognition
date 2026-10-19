# -*- coding: utf-8 -*-
"""Allergen database v2.1: apply the dietitian's decisions where they disagree with v2.0.

    python scripts/apply_expert_audit.py <filled form .xlsx>

The blind audit is scored on v2.0 (scripts/score_expert_audit.py) BEFORE this step. Afterwards,
for the audited dishes, every regulated allergen whose presence ('always'/'often' vs not) differs
between the dietitian and v2.0 takes the dietitian's level (source 'expert', version 2.1).
Updates data/allergen_db_v2/food_allergens.csv in place (run after finalize_allergen_db_v2.py build)
and writes migration_v2_0_to_v2_1.sql plus expert_audit/corrections_v2_1.csv.
"""
import csv
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from foodallergen import paths  # noqa: E402

D2 = paths.REPO_DATA / "allergen_db_v2"
PRESENT = {"Her zaman", "Sıklıkla"}


def main(form):
    ws = openpyxl.load_workbook(form, data_only=True)["Form"]
    hdr = [c.value for c in ws[1]]
    codes = [str(h).split("(")[-1].rstrip(")") if h and "(" in str(h) else None for h in hdr]
    sample = list(csv.DictReader(open(D2 / "expert_audit" / "sample.csv", encoding="utf-8")))
    foods = {r["YiyecekAdi"]: r["YiyecekID"] for r in csv.DictReader(open(D2 / "foods.csv", encoding="utf-8"))}
    aid = {r["Kod"]: r["AlerjenID"] for r in csv.DictReader(open(D2 / "allergens.csv", encoding="utf-8"))}
    rel_path = D2 / "food_allergens.csv"
    rel = list(csv.DictReader(open(rel_path, encoding="utf-8")))
    fields = list(rel[0])
    index = {(r["class_id"], r["code"]): r for r in rel}
    changes = []
    for row, s in zip(ws.iter_rows(min_row=2, values_only=True), sample):
        cid = s["class_id"]
        assert str(row[0]) == str(s["no"]), (row[0], s["no"])
        for v, c in zip(row, codes):
            if not c or c not in aid:
                continue
            expert = (v or "Yok").strip()
            cur = index.get((cid, c))
            db = cur["Miktar"] if cur else "Yok"
            if (expert in PRESENT) == (db in PRESENT):
                continue
            changes.append(dict(class_id=cid, code=c, v2_0=db, v2_1=expert, note=str(row[-1] or "")))
            if expert == "Yok":
                rel.remove(cur)
            elif cur:
                cur.update(Miktar=expert, Kaynak="expert", Surum="2.1")
            else:
                new = dict(YiyecekID=foods[cid], AlerjenID=aid[c], Miktar=expert, Kaynak="expert", Surum="2.1",
                           class_id=cid, code=c)
                rel.append(new)
                index[(cid, c)] = new
    with open(rel_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rel)
    with open(D2 / "expert_audit" / "corrections_v2_1.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(changes[0]))
        w.writeheader()
        w.writerows(changes)
    q = lambda s: "'" + str(s).replace("'", "''") + "'"  # noqa: E731
    sql = ["-- Allergen database v2.0 -> v2.1: corrections from the blind dietitian audit (run once).",
           "SET NAMES utf8mb4;", "START TRANSACTION;"]
    for ch in changes:
        y, a = foods[ch["class_id"]], aid[ch["code"]]
        if ch["v2_1"] == "Yok":
            sql.append(f"DELETE FROM yiyecekalerjenleri WHERE YiyecekID={y} AND AlerjenID={a};")
        elif ch["v2_0"] == "Yok":
            sql.append(f"INSERT INTO yiyecekalerjenleri (YiyecekID, AlerjenID, Miktar, Kaynak, Surum) VALUES "
                       f"({y}, {a}, {q(ch['v2_1'])}, 'expert', '2.1');")
        else:
            sql.append(f"UPDATE yiyecekalerjenleri SET Miktar={q(ch['v2_1'])}, Kaynak='expert', Surum='2.1' "
                       f"WHERE YiyecekID={y} AND AlerjenID={a};")
    sql += ["REPLACE INTO versiyonlar VALUES ('2.1', CURDATE(), 'v2.1: corrections from the blind dietitian audit (60 dishes)');",
            "COMMIT;",
            "SELECT COUNT(*) AS iliski FROM yiyecekalerjenleri;  -- beklenen: " + str(len(rel))]
    (D2 / "migration_v2_0_to_v2_1.sql").write_text("\n".join(sql) + "\n", encoding="utf-8")
    print(f"{len(changes)} corrections; relations now {len(rel)}")
    for ch in changes:
        print(ch["class_id"], ch["code"], ch["v2_0"], "->", ch["v2_1"])


if __name__ == "__main__":
    main(sys.argv[1])
