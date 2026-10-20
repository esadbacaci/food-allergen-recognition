# -*- coding: utf-8 -*-
"""Second-annotator check of the Turkish labels (reviewer 2: annotation consistency).

    python scripts/label_check_form.py form                 # 2 random images per Turkish class -> Excel
    python scripts/label_check_form.py score <filled.xlsx>  # agreement with the first annotator

The second annotator sees each photograph with the class name assigned by the first annotator and answers
whether the label is correct (Evet / Hayır / Emin değilim). Agreement = share of 'Evet' among definite answers.
Outputs: 01_veri/etiket_kontrol/{Etiket_Kontrol_Formu.xlsx, sample.csv}, 03_deneyler/analysis/label_check.json
"""
import csv
import json
import random
import sys
from pathlib import Path

from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from foodallergen import paths  # noqa: E402

OUT = paths.PROJECT / "01_veri" / "etiket_kontrol"
PER_CLASS = 2


def form():
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    cm = {c["class_id"]: c for c in csv.DictReader(open(paths.CLASS_MAP_V2, encoding="utf-8"))}
    man = [r for r in csv.DictReader(open(paths.MANIFEST_V2, encoding="utf-8"))
           if r["source"] != "food101" and cm[r["cls"]]["group"] != "food101"]
    rng = random.Random(2026)
    by = {}
    for r in man:
        by.setdefault(r["cls"], []).append(r)
    sample = [x for c in sorted(by) for x in rng.sample(by[c], min(PER_CLASS, len(by[c])))]
    rng.shuffle(sample)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "_thumbs").mkdir(exist_ok=True)
    with open(OUT / "sample.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["no", "cls", "path", "split"])
        w.writeheader()
        for i, r in enumerate(sample, 1):
            w.writerow(dict(no=i, cls=r["cls"], path=r["path"], split=r["split"]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kontrol"
    ws.append(["No", "Etiket (TR)", "Label (EN)", "Fotoğraf", "Etiket doğru mu? (Evet/Hayır/Emin değilim)", "Not"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col, wdt in zip("ABCDEF", [6, 24, 24, 30, 24, 30]):
        ws.column_dimensions[col].width = wdt
    dv = DataValidation(type="list", formula1='"Evet,Hayır,Emin değilim"', allow_blank=True)
    ws.add_data_validation(dv)
    img_root = paths.V2_IMAGES
    for i, r in enumerate(sample, start=2):
        c = cm[r["cls"]]
        ws.append([i - 1, c["display_tr"] or c["display_en"], c["display_en"], "", "", ""])
        t = OUT / "_thumbs" / f"{i - 1}.jpg"
        with Image.open(img_root / r["path"]) as im:
            im = im.convert("RGB")
            im.thumbnail((200, 200))
            im.save(t, quality=80)
        ws.add_image(XLImage(str(t)), f"D{i}")
        ws.row_dimensions[i].height = 155
        dv.add(f"E{i}")
    ws.freeze_panes = "A2"
    wb.save(OUT / "Etiket_Kontrol_Formu.xlsx")
    print(len(sample), "images ->", OUT / "Etiket_Kontrol_Formu.xlsx")


def score(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["Kontrol"]
    ans = [str(x[4] or "").strip().lower() for x in ws.iter_rows(min_row=2, values_only=True)]
    yes, no, unsure = ans.count("evet"), ans.count("hayır"), ans.count("emin değilim")
    res = dict(n=len(ans), yes=yes, no=no, unsure=unsure, unanswered=len(ans) - yes - no - unsure,
               agreement=yes / (yes + no) if yes + no else None)
    (paths.RUNS / "analysis" / "label_check.json").write_text(json.dumps(res, indent=1), encoding="utf-8")
    print(res)


if __name__ == "__main__":
    form() if sys.argv[1] == "form" else score(sys.argv[2])
